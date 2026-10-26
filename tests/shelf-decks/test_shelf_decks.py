"""Tests for the shelf-decks pack. Run: python3 -m unittest discover -s tests/shelf-decks"""
import io
import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'plugins/shelf-decks'))
from shelf_decks import brand, html_out, layout, review, sample, source  # noqa: E402
from shelf_decks.cli import main  # noqa: E402

try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False
try:
    import PIL  # noqa: F401
    HAVE_PIL = True
except ImportError:
    HAVE_PIL = False

BROWSERS = Path.home() / ('Library/Caches/ms-playwright' if sys.platform == 'darwin' else '.cache/ms-playwright')

DECK = """---
title: Test deck
---
[title]
# Hello Acme
## October 2026

---
# Points
- One **bold** point
  - a second-level point
- Two
notes: Say it slowly.

---
[stat]
# Numbers
- 38% :: More bookings
source: Fictional data, 2026

---
[bars]
# Channels
- Web :: 412
- Phone :: 186
source: Fictional data, 2026
"""


class Box(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='shelf-decks-test-')
        self.dir = Path(self.tmp.name)
        env = {'AIOS_CONFIG_DIR': str(self.dir / 'config'), 'HOME': str(self.dir)}
        if BROWSERS.is_dir() and 'PLAYWRIGHT_BROWSERS_PATH' not in os.environ:
            env['PLAYWRIGHT_BROWSERS_PATH'] = str(BROWSERS)  # keep using the installed Chromium under a fake HOME
        self.env = patch.dict(os.environ, env)
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def cli(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = main([str(a) for a in args])
        return code, out.getvalue(), err.getvalue()


class SourceTests(unittest.TestCase):
    def test_parse_layouts_and_fields(self):
        meta, slides = source.parse(DECK)
        self.assertEqual(meta['title'], 'Test deck')
        self.assertEqual([s['layout'] for s in slides], ['title', 'bullets', 'stat', 'bars'])
        self.assertEqual(slides[1]['bullets'], [('One **bold** point', 0), ('a second-level point', 1), ('Two', 0)])
        self.assertEqual(slides[1]['notes'], 'Say it slowly.')
        self.assertEqual(slides[2]['pairs'], [('38%', 'More bookings')])
        self.assertEqual([s['number'] for s in slides], [1, 2, 3, 4])

    def test_numbers_need_a_source(self):
        with self.assertRaises(source.DeckError) as cm:
            source.parse('# A\n\n---\n[stat]\n# B\n- 4 :: things\n')
        self.assertIn('source', str(cm.exception))

    def test_bars_need_numbers(self):
        with self.assertRaises(source.DeckError):
            source.parse('# A\n\n---\n[bars]\n# B\n- Web :: lots\n- Phone :: 3\nsource: x\n')

    def test_unknown_layout_and_missing_title(self):
        with self.assertRaises(source.DeckError):
            source.parse('[timeline]\n# A\n')
        with self.assertRaises(source.DeckError):
            source.parse('# A\n\n---\n- no title here\n')

    def test_two_column_needs_two_headings(self):
        with self.assertRaises(source.DeckError):
            source.parse('# A\n\n---\n[two-column]\n# B\n### Only one\n- x\n')

    def test_missing_image_refused(self):
        with self.assertRaises(source.DeckError):
            source.parse('# A\n\n---\n[image]\n# B\n![x](nope.png)\n', '/tmp')

    def test_word_count(self):
        _, slides = source.parse(DECK)
        self.assertEqual(source.words(slides[2]), 3)


class BrandTests(Box):
    def test_defaults_and_set(self):
        self.assertEqual(brand.load()['name'], 'Your Company')
        brand.set_value('colours.accent', '#C2410C')
        brand.set_value('name', 'Acme Studio')
        b = brand.load()
        self.assertEqual((b['name'], b['colours']['accent']), ('Acme Studio', '#C2410C'))

    def test_bad_colour_and_unknown_key(self):
        with self.assertRaises(brand.BrandError):
            brand.set_value('colours.accent', 'orange')
        with self.assertRaises(brand.BrandError):
            brand.set_value('colours.sparkle', '#000000')

    def test_svg_logo_refused(self):
        svg = self.dir / 'logo.svg'
        svg.write_text('<svg/>')
        with self.assertRaises(brand.BrandError):
            brand.set_value('logo', str(svg))

    def test_contrast_maths(self):
        self.assertEqual(brand.contrast('#000000', '#FFFFFF'), 21.0)
        self.assertEqual(brand.contrast('#FFFFFF', '#FFFFFF'), 1.0)
        report = {r['pair']: r['ok'] for r in brand.contrast_report(brand.load())}
        self.assertTrue(all(report.values()))

    def test_low_contrast_fails_review(self):
        brand.set_value('colours.muted', '#DDDDDD')
        bad = [r for r in brand.contrast_report(brand.load()) if not r['ok']]
        self.assertEqual(bad[0]['pair'], 'muted on background')


class LayoutTests(Box):
    def built(self):
        _, slides = source.parse(DECK)
        return slides, layout.build_all(slides, brand.load())

    def test_every_box_inside_the_slide(self):
        _, built = self.built()
        for s in built:
            for el in s['elements']:
                self.assertGreaterEqual(el['x'], 0)
                self.assertGreaterEqual(el['y'], 0)
                self.assertLessEqual(el['x'] + el['w'], layout.W)
                self.assertLessEqual(el['y'] + el['h'], layout.H)

    def test_content_slides_get_footer_and_number(self):
        _, built = self.built()
        footers = [e for e in built[1]['elements'] if e.get('role') == 'footer']
        self.assertEqual(len(footers), 2)
        self.assertEqual(footers[1]['paras'][0]['runs'][0][0], '2 / 4')
        self.assertFalse([e for e in built[0]['elements'] if e.get('role') == 'footer'])

    def test_bold_runs(self):
        self.assertEqual(layout.runs('One **bold** point'), [('One ', False), ('bold', True), (' point', False)])

    def test_bars_scale_to_largest(self):
        _, built = self.built()
        bars = [e for e in built[3]['elements'] if e['kind'] == 'rect' and e['h'] > 10]
        self.assertGreater(bars[0]['w'], bars[1]['w'])
        self.assertAlmostEqual(bars[1]['w'] / bars[0]['w'], 186 / 412, places=2)

    def test_html_is_self_contained(self):
        _, built = self.built()
        out = self.dir / 'deck.html'
        html_out.render({'title': 'T'}, built, brand.load(), out)
        text = out.read_text()
        self.assertIn('window.deckCheck', text)
        self.assertNotIn('src="http', text)
        self.assertEqual(text.count('<section class="s"'), 4)
        self.assertNotIn('font-family:\'Arial\',"', text, 'double quotes inside a style attribute break it')


@unittest.skipUnless(HAVE_PPTX and HAVE_PIL, 'python-pptx and Pillow needed')
class EndToEndTests(Box):
    def test_sample_build_review_export_verify(self):
        folder = self.dir / 'acme'
        self.assertEqual(self.cli('sample', folder)[0], 0)
        self.assertTrue((folder / 'acme-logo.png').exists())
        code, _, err = self.cli('export', folder)
        self.assertEqual(code, 1)
        self.assertIn('No HTML build yet', err)
        self.assertEqual(self.cli('build', folder)[0], 0)
        code, _, err = self.cli('export', folder)
        self.assertEqual(code, 1)
        self.assertIn('Review deck_v1 first', err)
        code, out, err = self.cli('review', folder)
        report = json.loads((folder / 'out/review_v1/review.json').read_text())
        if not report['measured']:
            self.skipTest('Chromium not available: ' + out)
        self.assertEqual(code, 0, out + err)
        self.assertEqual(len(report['images']), 10)
        self.assertEqual(self.cli('export', folder, '--pdf')[0], 0)
        self.assertTrue((folder / 'out/deck_v1.pdf').exists())
        code, out, _ = self.cli('verify', folder)
        self.assertEqual(code, 0, out)
        self.assertIn('Georgia', out)
        code, _, err = self.cli('export', folder)
        self.assertEqual(code, 1)
        self.assertIn('already exists', err)

    def test_edit_after_build_forces_rebuild(self):
        folder = self.dir / 'acme'
        self.cli('sample', folder)
        self.cli('build', folder)
        deck = folder / 'deck.md'
        deck.write_text(deck.read_text().replace('Q3 client review', 'Q3 review'))
        code, _, err = self.cli('review', folder)
        self.assertEqual(code, 1)
        self.assertIn('changed after deck_v1 was built', err)

    def test_overflow_is_caught(self):
        folder = self.dir / 'long'
        folder.mkdir()
        long = ' '.join(['word'] * 400)
        (folder / 'deck.md').write_text('# Long\n\n---\n# Too much\n' + long + '\n')
        self.cli('build', folder)
        code, out, _ = self.cli('review', folder)
        report = json.loads((folder / 'out/review_v1/review.json').read_text())
        if not report['measured']:
            self.skipTest('Chromium not available')
        self.assertEqual(code, 1)
        self.assertIn('does not fit', out)
        self.assertIn('words', out)

    def test_verify_catches_off_brand_font(self):
        from pptx import Presentation
        folder = self.dir / 'acme'
        self.cli('sample', folder)
        self.cli('build', folder)
        _, slides = source.parse((folder / 'deck.md').read_text(), folder)
        b = brand.load(folder / 'brand.json')
        target = folder / 'out/manual.pptx'
        from shelf_decks import pptx_out
        pptx_out.render({}, layout.build_all(slides, b), b, target)
        prs = Presentation(str(target))
        for shape in prs.slides[1].shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.name = 'Comic Sans MS'
        prs.save(str(target))
        result = review.verify(slides, b, target)
        self.assertFalse(result['ok'])
        self.assertIn('Comic Sans MS', ' '.join(result['problems']))


class SampleTests(Box):
    def test_sample_refuses_existing_deck(self):
        folder = self.dir / 'x'
        folder.mkdir()
        (folder / 'deck.md').write_text('# mine')
        with self.assertRaises(ValueError):
            sample.make(folder)

    def test_new_starter_parses(self):
        self.assertEqual(self.cli('new', self.dir / 'n', '--title', 'Plan')[0], 0)
        _, slides = source.parse((self.dir / 'n/deck.md').read_text())
        self.assertEqual(len(slides), 4)
        self.assertEqual(self.cli('new', self.dir / 'n')[0], 1)


if __name__ == '__main__':
    unittest.main()
