"""Write an editable PowerPoint from the same boxes as the HTML. Every word stays real, editable text."""
from pathlib import Path

from .brand import rgb
from .layout import H, W

EMU_PER_PX = 12192000 // W      # 6350: a 1920 px slide is 13.333 in wide
PT_PER_PX = 0.5


class PptxMissing(RuntimeError):
    pass


def _lib():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Emu, Pt
    except ImportError:
        raise PptxMissing('PowerPoint export needs python-pptx. Run: python3 -m pip install --user python-pptx')
    return locals()


def px(v):
    return int(round(v * EMU_PER_PX))


def _bullet(paragraph, size_px, level, colour, lib):
    qn = lib['qn']
    p_pr = paragraph._p.get_or_add_pPr()
    indent = px(size_px * 1.1)
    left = indent + px(size_px * 1.2) * level
    p_pr.set('marL', str(left))
    p_pr.set('indent', str(-indent))
    for tag in ('a:buClr', 'a:buFont', 'a:buChar', 'a:buNone'):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    clr = p_pr.makeelement(qn('a:buClr'), {})
    srgb = clr.makeelement(qn('a:srgbClr'), {'val': colour.lstrip('#').upper()})
    clr.append(srgb)
    p_pr.append(clr)
    p_pr.append(p_pr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
    p_pr.append(p_pr.makeelement(qn('a:buChar'), {'char': '•'}))


def _text(slide, el, brand, lib):
    Pt, RGBColor = lib['Pt'], lib['RGBColor']
    box = slide.shapes.add_textbox(px(el['x']), px(el['y']), px(el['w']), px(el['h']))
    box.name = 'shelf-%s' % el['role']
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = lib['MSO_AUTO_SIZE'].NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'top': lib['MSO_ANCHOR'].TOP, 'middle': lib['MSO_ANCHOR'].MIDDLE,
                          'bottom': lib['MSO_ANCHOR'].BOTTOM}[el['anchor']]
    align = {'left': lib['PP_ALIGN'].LEFT, 'center': lib['PP_ALIGN'].CENTER, 'right': lib['PP_ALIGN'].RIGHT}[el['align']]
    font_name = brand['fonts'][el['font']]
    for i, p in enumerate(el['paras']):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = el['line']
        if i < len(el['paras']) - 1 and el['space']:
            paragraph.space_after = Pt(el['space'] * PT_PER_PX)
        if p['bullet']:
            _bullet(paragraph, el['size'], p['level'], brand['colours']['accent'], lib)
        for t, bold in p['runs']:
            run = paragraph.add_run()
            run.text = t
            f = run.font
            f.name = font_name
            f.size = Pt(el['size'] * PT_PER_PX)
            f.bold = bool(el['bold'] or bold)
            f.italic = bool(el['italic'])
            f.color.rgb = RGBColor(*rgb(el['colour']))
    return box


def render(meta, built, brand, out_path):
    lib = _lib()
    prs = lib['Presentation']()
    prs.slide_width, prs.slide_height = lib['Emu'](px(W)), lib['Emu'](px(H))
    blank = prs.slide_layouts[6]
    for s in built:
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = lib['RGBColor'](*rgb(s['background']))
        for el in s['elements']:
            if el['kind'] == 'rect':
                shape = slide.shapes.add_shape(lib['MSO_SHAPE'].RECTANGLE, px(el['x']), px(el['y']), px(el['w']), px(el['h']))
                shape.fill.solid()
                shape.fill.fore_color.rgb = lib['RGBColor'](*rgb(el['fill']))
                shape.line.fill.background()
                shape.shadow.inherit = False
                style = shape._element.find(lib['qn']('p:style'))
                if style is not None:  # the theme style adds a shadow and an outline in some apps
                    shape._element.remove(style)
                shape.name = 'shelf-rect'
            elif el['kind'] == 'image':
                pic = slide.shapes.add_picture(el['path'], px(el['x']), px(el['y']), px(el['w']), px(el['h']))
                pic.name = 'shelf-image'
            else:
                _text(slide, el, brand, lib)
        if s['notes']:
            slide.notes_slide.notes_text_frame.text = s['notes']
    props = prs.core_properties
    props.title = meta.get('title') or (built[0]['title'] if built else '')
    props.author = meta.get('author') or brand.get('name', '')
    props.comments = 'Made with shelf-decks 2.0.0'
    prs.save(str(out_path))
    return out_path


def inspect(path):
    """Read a PPTX back: what a reviewer needs to know without opening PowerPoint."""
    lib = _lib()
    prs = lib['Presentation'](str(path))
    slides = []
    for n, slide in enumerate(prs.slides, 1):
        fonts, colours, texts, images = set(), set(), [], 0
        for shape in slide.shapes:
            if shape.shape_type == 13:
                images += 1
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            fonts.add(r.font.name)
                        try:
                            colours.add('#' + str(r.font.color.rgb))
                        except AttributeError:
                            pass
                    if p.text.strip():
                        texts.append(p.text)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
        slides.append({'number': n, 'fonts': sorted(fonts), 'colours': sorted(colours), 'texts': texts,
                       'images': images, 'notes': bool(notes.strip())})
    return {'width_px': round(prs.slide_width / EMU_PER_PX), 'height_px': round(prs.slide_height / EMU_PER_PX),
            'slides': slides}


def exists(path):
    return Path(path).exists()
